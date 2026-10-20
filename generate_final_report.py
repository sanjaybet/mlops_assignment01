from io import BytesIO
import os
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT_DIR = Path(__file__).resolve().parent
VIDEO_URL = 'https://drive.google.com/file/d/149757pNEOqkvSM3HhpsGu9iQHXtvTMpI/view?usp=sharing'
REPOSITORY_URL = 'https://github.com/sanjaybet/mlops_assignment01'
ACTION_RUN_URL = 'https://github.com/sanjaybet/mlops_assignment01/actions/runs/39002949584'


def add_section_heading(document, title):
    document.add_heading(title, level=1)


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True,
    )
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relationship_id)
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle')
    style.set(qn('w:val'), 'Hyperlink')
    properties.append(style)
    run.append(properties)
    text_element = OxmlElement('w:t')
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_table_of_contents(document):
    document.add_heading('Table of Contents', level=1)
    paragraph = document.add_paragraph()
    run = paragraph.add_run()

    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    placeholder = OxmlElement('w:t')
    placeholder.text = 'Open in Word and update the field to populate the table of contents.'

    for field_type in ('begin', 'separate'):
        field_char = OxmlElement('w:fldChar')
        field_char.set(qn('w:fldCharType'), field_type)
        run._r.append(field_char)
    run._r.append(instruction)
    run._r.append(placeholder)
    field_end = OxmlElement('w:fldChar')
    field_end.set(qn('w:fldCharType'), 'end')
    run._r.append(field_end)


def add_image_with_caption(document, image_stream, caption):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_source = str(image_stream) if isinstance(image_stream, Path) else image_stream
    paragraph.add_run().add_picture(image_source, width=Inches(6.25))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_landscape_screenshot_pages(document, screenshots):
    for image_path, caption in screenshots:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Inches(0.08)
        paragraph.add_run().add_picture(
            str(image_path),
            width=Inches(6.25),
            height=Inches(3.11),
        )
        caption_paragraph = document.add_paragraph(caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.keep_together = True
        caption_paragraph.paragraph_format.space_before = Inches(0.05)
        caption_paragraph.paragraph_format.space_after = Inches(0)

    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)


def render_architecture_diagram():
    figure, axis = plt.subplots(figsize=(12, 6.8), dpi=180)
    figure.patch.set_facecolor('#f4f6f4')
    axis.set_facecolor('#f4f6f4')
    axis.set_xlim(0, 12)
    axis.set_ylim(0, 7)
    axis.axis('off')

    nodes = {
        'data': (0.5, 5.1, 'UCI Cleveland\nHeart Disease data'),
        'ingest': (3.2, 5.1, 'Download +\nschema validation'),
        'eda': (5.9, 5.1, 'EDA + cleaning\nprocessed data'),
        'train': (8.6, 5.1, 'Model training\n5-fold CV + tuning'),
        'mlflow': (8.6, 3.1, 'MLflow tracking\nparams, metrics, artifacts'),
        'model': (5.9, 3.1, 'Champion pipeline\nmodel.joblib'),
        'api': (3.2, 3.1, 'FastAPI\n/health + /predict'),
        'docker': (0.5, 3.1, 'Docker image\nPython 3.12'),
        'k8s': (0.5, 0.8, 'Kubernetes\nDeployment + Service'),
        'ci': (5.9, 0.8, 'GitHub Actions\nlint, tests, EDA, training'),
    }
    box_width, box_height = 2.25, 1.0
    colors = {
        'data': '#dce9e3', 'ingest': '#dce9e3', 'eda': '#dce9e3',
        'train': '#f2dfba', 'mlflow': '#d9e6ef', 'model': '#f2dfba',
        'api': '#d9e6ef', 'docker': '#e9e1d2', 'k8s': '#e9e1d2',
        'ci': '#e8d9ca',
    }
    for key, (x, y, label) in nodes.items():
        box = FancyBboxPatch(
            (x, y), box_width, box_height,
            boxstyle='round,pad=0.12,rounding_size=0.12',
            linewidth=1.2, edgecolor='#263b3a', facecolor=colors[key],
        )
        axis.add_patch(box)
        axis.text(
            x + box_width / 2, y + box_height / 2, label,
            ha='center', va='center', fontsize=9, color='#1c2c2a',
            fontweight='medium',
        )

    arrows = [
        ((2.8, 5.6), (3.15, 5.6)),
        ((5.5, 5.6), (5.85, 5.6)),
        ((8.2, 5.6), (8.55, 5.6)),
        ((9.7, 5.0), (7.2, 4.25)),
        ((5.85, 3.6), (5.5, 3.6)),
        ((3.15, 3.6), (2.8, 3.6)),
        ((1.6, 3.0), (1.6, 1.9)),
        ((9.7, 5.0), (9.7, 4.25)),
        ((7.05, 1.9), (9.1, 5.0)),
    ]
    for start, end in arrows:
        axis.annotate(
            '', xy=end, xytext=start,
            arrowprops={'arrowstyle': '-|>', 'lw': 1.4, 'color': '#496563'},
        )

    axis.text(
        6, 6.55, 'HEART DISEASE PREDICTION | END-TO-END MLOPS ARCHITECTURE',
        ha='center', va='center', fontsize=12, fontweight='bold',
        color='#203431',
    )
    image_stream = BytesIO()
    figure.savefig(image_stream, format='png', bbox_inches='tight', facecolor=figure.get_facecolor())
    plt.close(figure)
    image_stream.seek(0)
    return image_stream


def render_source_capture(title, relative_path, start_line=1, end_line=None):
    source_path = PROJECT_DIR / relative_path
    source_lines = source_path.read_text(encoding='utf-8').splitlines()
    selected_lines = source_lines[start_line - 1:end_line]
    figure_height = max(2.5, 1.1 + 0.19 * len(selected_lines))
    figure, axis = plt.subplots(figsize=(6.5, figure_height), dpi=180)
    figure.patch.set_facecolor('#172126')
    axis.set_facecolor('#172126')
    axis.axis('off')
    axis.text(
        0.035, 0.96, f'{title}  |  {relative_path}',
        transform=axis.transAxes, ha='left', va='top',
        fontsize=9, fontweight='bold', color='#f2c879',
    )
    code = '\n'.join(
        f'{line_number:02d}  {line}'
        for line_number, line in enumerate(selected_lines, start=start_line)
    )
    axis.text(
        0.035, 0.87, code, transform=axis.transAxes,
        ha='left', va='top', fontsize=7.1, color='#e3ece8',
        fontfamily='DejaVu Sans Mono', linespacing=1.35,
    )
    image_stream = BytesIO()
    figure.savefig(image_stream, format='png', bbox_inches='tight', facecolor=figure.get_facecolor())
    plt.close(figure)
    image_stream.seek(0)
    return image_stream


document = Document()
document.styles['Normal'].font.name = 'Calibri'
document.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
update_fields = OxmlElement('w:updateFields')
update_fields.set(qn('w:val'), 'true')
document.settings.element.append(update_fields)

paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_before = Pt(105)
paragraph.paragraph_format.space_after = Pt(10)
run = paragraph.add_run('MLOps (S1-26_AIMLZG523) Assignment, 2026')
run.bold = True
run.font.size = Pt(22)

paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_after = Pt(34)
run = paragraph.add_run('Heart Disease Prediction Pipeline')
run.bold = True
run.font.size = Pt(16)

for label, value in [
    ('Student Id: ', '2025ae05368'),
    ('Student Name: ', 'Sanjay. S. Betageri'),
]:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.add_run(label).bold = True
    paragraph.add_run(value)

repository_paragraph = document.add_paragraph()
repository_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
repository_paragraph.add_run('Repository: ').bold = True
add_hyperlink(repository_paragraph, REPOSITORY_URL, REPOSITORY_URL)
video_paragraph = document.add_paragraph()
video_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
video_paragraph.add_run('Video demo link: ').bold = True
add_hyperlink(video_paragraph, VIDEO_URL, VIDEO_URL)

paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_before = Pt(26)
paragraph.paragraph_format.space_after = Pt(6)
run = paragraph.add_run('Environment: Prayogshala and local setup')
run.bold = True
run.font.size = Pt(13)

paragraph = document.add_paragraph(
    'Deployment validation was performed in the Prayogshala virtual lab; '
    'the project can also be set up and run locally using the documented '
    'Python environment.'
)
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.left_indent = Inches(0.7)
paragraph.paragraph_format.right_indent = Inches(0.7)
paragraph.paragraph_format.space_after = Pt(0)
document.add_page_break()

add_table_of_contents(document)
document.add_page_break()

add_section_heading(document, '1. Introduction and objective')
document.add_paragraph('This report presents a reproducible heart disease prediction solution using the UCI Cleveland dataset. Its MLOps workflow covers data ingestion and cleaning, model development and tracking, API serving, deployment, testing, and automation.')
document.add_paragraph('The objective was to build a pipeline that can be recreated from a clean environment, validated automatically, and served through a prediction API using Python, scikit-learn, MLflow, FastAPI, Docker, Kubernetes, and GitHub Actions.')

add_section_heading(document, '2. Environment and setup')
document.add_paragraph('The project uses an isolated Python 3.12 virtual environment. Dependencies for data processing, training, MLflow, FastAPI, testing, and report generation are listed in requirements.txt; setup commands are documented in README.md.')
document.add_paragraph('The requirements pin setuptools below version 81 to avoid an MLflow compatibility conflict. Recreating the environment from the requirements file supports consistent local and CI runs.')
document.add_heading('Steps to run the code', level=2)
run_commands = (
    'cd ~/workspace/mlops_assignment01\n\n'
    '# Create and activate the environment\n'
    'python -m venv .venv\n'
    'source .venv/bin/activate\n\n'
    '# Install project dependencies\n'
    'python -m pip install --upgrade pip\n'
    'pip install -r requirements.txt\n\n'
    '# Download data, clean it, then train and log experiments\n'
    'python src/download_data.py\n'
    'python src/eda.py\n'
    'python src/train.py\n\n'
    '# Start the prediction API\n'
    'uvicorn app.main:app --host 0.0.0.0 --port 8001'
)
commands_paragraph = document.add_paragraph()
commands_paragraph.paragraph_format.left_indent = Inches(0.15)
commands_paragraph.paragraph_format.space_before = Pt(2)
commands_paragraph.paragraph_format.space_after = Pt(4)
commands_run = commands_paragraph.add_run(run_commands)
commands_run.font.name = 'Consolas'
commands_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
commands_run.font.size = Pt(8)

add_section_heading(document, '3. Data acquisition and validation')
document.add_paragraph('The acquisition script downloads the official UCI Heart Disease archive, extracts the Cleveland data, and validates its expected schema and records before use. This makes the raw input traceable and causes malformed or incomplete data to fail early.')
document.add_paragraph('The raw and processed data, download script, and run instructions are included in the repository.')

add_section_heading(document, '4. Exploratory data analysis and preprocessing')
document.add_paragraph('The EDA workflow validates the schema and target, removes duplicate records, and maps target 0 to “no disease” and targets 1–4 to “disease present.” It saves the cleaned dataset for training and generates feature-distribution, class-balance, and correlation plots.')
document.add_paragraph('Missing feature values are retained for fold-local imputation in the training pipeline. Numeric features are imputed and scaled; categorical features are imputed and one-hot encoded, preventing preprocessing leakage across validation folds.')

add_section_heading(document, '5. Model development and evaluation')
document.add_paragraph('Logistic Regression and Random Forest are compared using stratified five-fold cross-validation and a holdout test set. Preprocessing (imputation, scaling, and encoding) is fitted within each pipeline fold to prevent data leakage. Hyperparameters are selected by cross-validated ROC-AUC; the holdout set is reserved for final reporting.')
document.add_paragraph('Confusion matrices and ROC curves are saved for both candidates to support review of their classification performance.')
document.add_heading('End-to-end architecture', level=2)
document.add_paragraph('The architecture connects validated UCI data to preprocessing and model selection, records experiment evidence in MLflow, and packages the selected inference pipeline behind FastAPI. GitHub Actions automates linting, tests, preprocessing, and training; Docker and Kubernetes provide the deployment path.')
add_image_with_caption(
    document,
    render_architecture_diagram(),
    'Figure 1. End-to-end data, training, tracking, and serving architecture.',
)
document.add_page_break()

add_section_heading(document, '6. MLflow experiment tracking')
document.add_paragraph('MLflow records each candidate’s parameters, preprocessing choices, cross-validation and holdout metrics, tuning results, plots, configuration, and fitted model. This preserves the evidence behind model selection and supports repeatable comparison.')
document.add_paragraph('The recorded runs show the following ROC-AUC comparison:')
tracking_table = document.add_table(rows=1, cols=3)
tracking_table.style = 'Light Shading Accent 1'
for cell, heading in zip(
    tracking_table.rows[0].cells,
    ['Candidate model', 'CV ROC-AUC (selection)', 'Holdout ROC-AUC (reporting)'],
):
    cell.text = heading
for model_name, cv_auc, test_auc in [
    ('Logistic Regression', '0.9025', '0.9665'),
    ('Random Forest', '0.9030', '0.9448'),
]:
    cells = tracking_table.add_row().cells
    cells[0].text = model_name
    cells[1].text = cv_auc
    cells[2].text = test_auc
document.add_paragraph('Random Forest is the champion by CV ROC-AUC (0.9030 vs. 0.9025). Holdout ROC-AUC is reported independently and does not determine selection. The selected pipeline is saved to data/processed/model.joblib for inference.')

add_section_heading(document, '7. Deployment of the prediction service')
document.add_paragraph('Once the model was trained and validated, it was exposed through a FastAPI application. The prediction service receives patient information in JSON format, loads the serialized model artifact, and returns the disease prediction together with probability scores for both disease and no-disease classes. This API layer is the operational entry point for downstream users and stakeholders.')
document.add_paragraph('The API implementation includes health-check endpoints and an inference endpoint. This design is useful for both local testing and deployment monitoring because it allows the service to be easily validated by a lightweight HTTP request. The service is designed to be stateless and portable, which makes it suitable for containerized deployment.')
document.add_paragraph('The API was deployed to a local Minikube cluster running with the Docker driver in the Prayogshala virtual lab. The deployment reached two ready API replicas and one ready Prometheus replica. The API Service is Kubernetes type LoadBalancer; Minikube did not assign an external IP, so verification used kubectl port-forward, which is the documented local access path for this cluster.')
document.add_paragraph('Runtime checks returned HTTP 200 from /health and /predict. The prediction response contained a binary class and probabilities for both classes. The /metrics endpoint exposed request counters, and a Prometheus query for up{job="heart-disease-api"} returned 1 for the API target. These screenshots document the observed lab run; they are evidence of local Kubernetes deployment, not a public cloud endpoint.')
add_landscape_screenshot_pages(
    document,
    [
        (
            PROJECT_DIR / 'screenshots' / 'Prayogshala-docker-kubernetes-1.png',
            'Figure 8. Prayogshala Minikube deployment: API and Prometheus pods are ready; the local LoadBalancer has no external IP.',
        ),
        (
            PROJECT_DIR / 'screenshots' / 'Prayogshala-kubernetes-API-Call.png',
            'Figure 9. Successful health and prediction requests, with Prometheus-format API request counters.',
        ),
    ],
)

add_section_heading(document, '8. CI/CD and Dockerization')
document.add_paragraph('The project uses GitHub Actions for automated validation. The workflow checks out the repository, installs the requirements, runs linting, executes tests, and then triggers the data processing and training pipeline. This is a standard continuous integration pattern and helps enforce quality standards on every change to the repository.')
document.add_paragraph('The pipeline is designed to fail clearly if any component of the workflow fails. This is important because production systems should not silently accept invalid code or broken environment assumptions. In the project, the workflow therefore acts as a guardrail by running the tests and model training automatically before the code is considered acceptable.')
document.add_paragraph('The application is also containerized using Docker. The Dockerfile creates a slim Python environment, installs the project requirements, copies the source code, and starts the FastAPI server. This allows the project to be run consistently across local development, CI, and deployment environments. The container architecture is an essential MLOps practice because it minimizes environment drift and simplifies portability.')
document.add_paragraph('Together, the GitHub Actions workflow and Docker configuration ensure both automation and reproducibility. This covers the assignment requirement for pipeline validation, deployment packaging, and environment consistency.')
document.add_paragraph('GitHub Actions run #8 completed successfully on the main branch after pull request #4 was merged. Commit 98dfa02 passed the lint, test, and model-training steps in 57 seconds. The first two figures below are screenshots of the hosted run summary and its completed job steps. The Dockerfile and Kubernetes configuration captures document the checked-in deployment setup; the Minikube runtime screenshots in Section 7 separately show the lab deployment.')
document.add_heading('Successful hosted GitHub Actions run', level=2)
add_image_with_caption(
    document,
    PROJECT_DIR / 'screenshots' / 'CI-1.png',
    'Figure 2. GitHub Actions run #8 succeeded on main at commit 98dfa02 (57 seconds).',
)
add_image_with_caption(
    document,
    PROJECT_DIR / 'screenshots' / 'CI-2.png',
    'Figure 3. The lint-and-test job completed successfully, including lint, tests, and model training.',
)
run_details = document.add_paragraph('Run details: ')
add_hyperlink(run_details, ACTION_RUN_URL, ACTION_RUN_URL)
document.add_heading('GitHub Actions workflow configuration', level=2)
add_image_with_caption(
    document,
    render_source_capture(
        'CI workflow configuration', '.github/workflows/cicd.yaml', 1, 30
    ),
    'Figure 4. Source capture of the CI workflow: Python setup, lint, tests, EDA, and training.',
)
document.add_heading('Container image configuration', level=2)
add_image_with_caption(
    document,
    render_source_capture('Docker image configuration', 'Dockerfile', 1, 20),
    'Figure 5. Source capture of the Dockerfile used to package the FastAPI service.',
)
document.add_heading('Kubernetes deployment configuration', level=2)
add_image_with_caption(
    document,
    render_source_capture(
        'Kubernetes Deployment configuration', 'k8s/deployment.yaml', 1, 40
    ),
    'Figure 6. Source capture of the Kubernetes Deployment, replicas, ports, and health probes.',
)
add_image_with_caption(
    document,
    render_source_capture('Kubernetes Service configuration', 'k8s/service.yaml', 1, 30),
    'Figure 7. Source capture of the LoadBalancer Service routing traffic to the API pods.',
)
document.add_page_break()

add_section_heading(document, '9. Monitoring, logging, and runtime evidence')
document.add_paragraph('The FastAPI application logs request method, path, status code, and duration to the container output stream without recording patient request bodies. Kubernetes captures these messages through pod logs. The Prometheus FastAPI instrumentator exposes HTTP request metrics at /metrics; the in-cluster Prometheus Deployment scrapes the API Service every 10 seconds.')
document.add_paragraph('During the Prayogshala run, the API metrics output included successful request counters for /health, /metrics, and /predict. Prometheus returned a value of 1 for up{job="heart-disease-api"}, confirming that the configured target was reachable and being scraped. The first Prometheus capture shows the port-forward used to open its UI; the second shows the successful query and target value.')
add_landscape_screenshot_pages(
    document,
    [
        (
            PROJECT_DIR / 'screenshots' / 'Prometheus-1.png',
            'Figure 10. Port-forward from the Prayogshala Minikube cluster to the Prometheus UI.',
        ),
        (
            PROJECT_DIR / 'screenshots' / 'Prometheus-2.png',
            'Figure 11. Prometheus query up{job="heart-disease-api"} returns 1 for the deployed API target.',
        ),
    ],
)

add_section_heading(document, '10. Production-readiness and compliance check')
document.add_paragraph('The container was deployed to Prayogshala Minikube. Health, prediction, and metrics endpoints were checked, and Prometheus confirmed the API target was up. CI runs linting, tests, EDA, and training. The local cluster is session-scoped; the repository contains the manifests and instructions to recreate it.')
document.add_paragraph('The repository includes the source, scripts, model/data artifacts, tests, deployment manifests, screenshots, and report required for review.')

add_section_heading(document, '11. Documentation and reporting')
document.add_paragraph('This report summarizes setup, data processing, model selection, tracking, deployment, CI/CD, and validation evidence. The Word report and supporting repository documentation are included for assessment.')
video_paragraph = document.add_paragraph('Video demo link: ')
add_hyperlink(video_paragraph, VIDEO_URL, VIDEO_URL)
repository_paragraph = document.add_paragraph('Repository link: ')
add_hyperlink(repository_paragraph, REPOSITORY_URL, REPOSITORY_URL)

add_section_heading(document, '12. Deliverables checklist')
for item in [
    'GitHub repository with source code and model artifacts',
    'Dockerfile and dependency file',
    'Dataset download and preprocessing scripts',
    'EDA, training, and inference scripts',
    'Unit tests under tests/',
    'GitHub Actions workflow YAML',
    'Deployment manifests for Kubernetes',
    'Screenshot folder with visual outputs',
    'Final written report in Word/PDF format',
    'Local access instructions or public API URL',
]:
    document.add_paragraph(f'• {item}')

video_item = document.add_paragraph('• Video demo: ')
add_hyperlink(video_item, VIDEO_URL, VIDEO_URL)

document.add_paragraph('')
document.add_paragraph('Conclusion: The project was implemented and validated as a complete MLOps pipeline for heart disease prediction. It covers data acquisition, cleansing, experimentation, deployment, automation, and reporting in a way that aligns with the assignment specification and production-readiness expectations.')

output_path = os.environ.get(
    'REPORT_OUTPUT_PATH', '2025ae05368_mlops_assignment1_report.docx'
)
if __name__ == '__main__':
    document.save(output_path)
    print(f'Created {output_path}')
